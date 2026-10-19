"""The private store: local, unpublished genomes that belong in a final analysis dataset but must never mix with the public curated set.

Why a store of its own (and not rows in `genome`):
* the curated set is defined by NCBI (INSDC) accessions - a private genome would be reported as "gone from NCBI" by the next check;
* the public page, the Superset view, TSV downloads and backups are all built from `genome`: keeping private data out of that table keeps it
  out of every one of them by construction.

How genomes get in
* a **spreadsheet** (.xlsx, or .tsv / .csv) with the columns the Nextstrain mpox workflow uses (`strain`, `date`, `country`, `division`,
  `location`, `host`, `clade`, ...) and an id column `PPX_accession` or `accession`, plus a FASTA whose ids are those ids; `template_xlsx()` makes
  a ready-to-fill workbook (a sheet for the data, one explaining every column, one with an example row);
* **one genome at a time** (`one()`): the fields of a single row and one sequence (a file or pasted text).
Both go through the same dry run (`check`): structural problems are errors that stop the import (missing / duplicate / colliding ids, rows
without a sequence, sequences without a row, bad sequence characters, a genome below the coverage floor); format quirks (partial dates, a missing
region, a length that differs) are normalised and reported as warnings. `defaults` fills blank cells (e.g. clade = Ib for a clade Ib batch).

**Every upload is a version.** An id that is uploaded again with changes becomes version 2, 3, ... with what changed (which fields, whether the
sequence); uploading the same content again changes nothing. Old versions (and their sequences) are kept - also after a delete - and any version
can be restored (as a new version). Sequences live under `data/private/<virus>/`, old ones under `_versions/`.

QC (sequence statistics and Nextclade) runs on the current version under the accession `private:<id>`. Private rows leave the tool only through an
export that asks for them explicitly (`export --with-private`), as separate files (`metadata-private.tsv`, `sequences-private.fasta`)."""
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path

from . import completeness, config, export, log
from .db import audit, connect

_log = log.get("private")
ID_COLUMNS = ("PPX_accession", "accession")
REQUIRED = ("strain", "date")
ACGT = set("ACGTRYSWKMBDHVN-")        # IUPAC nucleotide codes (and gaps); anything else is not a nucleotide sequence
RESERVED = {"qc", "source_nextstrain", "source_private"}
SHEET = "genomes"

# the columns of the template: (name, required, meaning, example, allowed values or None)
COLUMNS = [
    ("PPX_accession", True, "이 유전체의 고유 id (FASTA 헤더와 같아야 함). 영문·숫자·`._-`만, 공개 accession과 겹치면 안 됨", "LAB_0001", None),
    ("strain", True, "균주/검체 이름", "MPXV/Uganda/KLA-0001/2025", None),
    ("date", True, "수집일: YYYY-MM-DD, 모르는 부분은 XX (2025-03-XX). 2025-03 처럼 써도 자동 변환", "2025-03-14", None),
    ("country", True, "국가 (영문 표준 이름)", "Uganda", None),
    ("division", False, "도/주 등 1차 행정구역", "Kampala", None),
    ("location", False, "시/군 등 더 작은 지역", "Nakawa", None),
    ("host", False, "숙주 학명", "Homo sapiens", None),
    ("clade", False, "clade (비우면 업로드 기본값 → Nextclade 순). Nextclade와 다르면 표시됩니다", "Ib", None),
    ("lineage", False, "lineage (비우면 Nextclade 값)", "A.4.1", None),
    ("outbreak", False, "outbreak (비우면 Nextclade 값)", "sh2023", None),
    ("institution", False, "제출 기관", "Makerere University", None),
    ("authors", False, "저자 (예: Kato et al.)", "Kato et al.", None),
    ("date_released", False, "공개일/보관일 YYYY-MM-DD", "", None),
]
CLADE_CHOICES = {"orthopoxvirus monkeypox": ["Ia", "Ib", "IIa", "IIb"]}


def prefix(pid: str) -> str:
    return f"private:{pid}"


def private_dir(virus: str) -> Path:
    return config.DATA_DIR / "private" / export.slug(virus)


def _safe(pid: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", pid)


# ------------------------------------------------------------------------------------------------ reading input
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_table(data, filename: str = "") -> tuple[list[str], list[dict]]:
    """A spreadsheet (.xlsx), or tab / comma separated text. -> (columns, rows of strings); rows that are entirely empty are skipped."""
    raw = data if isinstance(data, (bytes, bytearray)) else str(data).encode("utf-8")
    if filename.lower().endswith(".xlsx") or raw[:2] == b"PK":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("스프레드시트(.xlsx)를 읽으려면 openpyxl이 필요합니다 (`mamba install openpyxl`). TSV/CSV는 그대로 쓸 수 있습니다.") from None
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"스프레드시트를 열 수 없습니다: {e}") from None
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if not head:
            return [], []
        cols = [_cell(c) for c in head]
        rows = []
        for r in it:
            vals = [_cell(c) for c in r]
            if any(vals):
                rows.append({cols[i]: (vals[i] if i < len(vals) else "") for i in range(len(cols)) if cols[i]})
        return [c for c in cols if c], rows
    text = raw.decode("utf-8-sig", "replace")
    first = text.splitlines()[0] if text.strip() else ""
    delim = "\t" if "\t" in first else ("," if "," in first else "\t")
    rd = csv.DictReader(io.StringIO(text), delimiter=delim)
    cols = [c.strip() for c in (rd.fieldnames or []) if c]
    rows = []
    for r in rd:
        row = {(k or "").strip(): (v or "").strip() for k, v in r.items() if k is not None}
        if any(row.values()):
            rows.append(row)
    return cols, rows


def _read_fasta(text: str) -> tuple[dict[str, str], list[str]]:
    seqs: dict[str, str] = {}
    dups: list[str] = []
    cur, buf = None, []
    for ln in text.splitlines():
        if ln.startswith(">"):
            if cur is not None:
                (dups.append(cur) if cur in seqs else None)
                seqs[cur] = "".join(buf).upper()
            cur, buf = ln[1:].split()[0] if ln[1:].strip() else "", []
        elif cur is not None:
            buf.append(ln.strip())
    if cur is not None:
        (dups.append(cur) if cur in seqs else None)
        seqs[cur] = "".join(buf).upper()
    return seqs, dups


def _public_ids(con, watch_id: int) -> set[str]:
    ids = set()
    for r in con.execute("SELECT accession, accession_version, pathoplexus_accession FROM genome WHERE watch_id=?", (watch_id,)):
        ids.add(r[0])
        if r[1]:
            ids.add(f"{r[0]}.{r[1]}")
        if r[2]:
            ids.add(str(r[2]).split(".")[0])
    return ids


# ------------------------------------------------------------------------------------------------ dry run
def check(watch_id: int, metadata, fasta_text: str, id_column: str | None = None, defaults: dict | None = None, filename: str = "") -> dict:
    """Dry run: parse, validate and normalise. Nothing is written.
    -> {ok, errors, warnings, rows (normalised), n, new, update, unchanged, id_column}"""
    con = connect()
    try:
        w = con.execute("SELECT name FROM watchlist WHERE id=?", (watch_id,)).fetchone()
        if not w:
            raise ValueError("바이러스를 찾을 수 없습니다.")
        public = _public_ids(con, watch_id)
        have = {r[0]: r for r in con.execute("SELECT id, metadata_json, seq_md5 FROM private_genome WHERE watch_id=?", (watch_id,))}
        ref = completeness.reference_map(con).get(watch_id)
    finally:
        con.close()
    errors, warnings = [], []
    cols, rows = _read_table(metadata, filename)
    seqs, dup_seq = _read_fasta(fasta_text)
    idc = id_column or next((c for c in ID_COLUMNS if c in cols), None)
    defaults = {k: v for k, v in (defaults or {}).items() if v}
    if not rows:
        errors.append("메타데이터에 행이 없습니다.")
    if not idc or idc not in cols:
        errors.append(f"id 열이 없습니다. 메타데이터에 {' 또는 '.join(ID_COLUMNS)} 열이 있어야 합니다 (FASTA 헤더와 같은 값).")
    for need in REQUIRED:
        if need not in cols:
            errors.append(f"필수 열 `{need}`가 없습니다.")
    if errors:
        return {"ok": False, "errors": errors, "warnings": warnings, "rows": [], "n": 0, "new": 0, "update": 0, "unchanged": 0, "id_column": idc}
    seen: dict[str, int] = {}
    out = []
    for i, r in enumerate(rows, start=2):                         # line numbers as in the file (header = line 1)
        pid = r.get(idc, "")
        if not pid:
            errors.append(f"{i}행: id가 비어 있습니다.")
            continue
        if not re.fullmatch(r"[A-Za-z0-9._-]+", pid):
            errors.append(f"{i}행: id `{pid}`에 허용되지 않는 문자가 있습니다 (영문·숫자·`._-`만).")
            continue
        if pid in seen:
            errors.append(f"{i}행: id `{pid}`가 {seen[pid]}행과 겹칩니다.")
            continue
        seen[pid] = i
        if pid in public:
            errors.append(f"{i}행: id `{pid}`는 공개 데이터(NCBI/Pathoplexus)의 accession과 같습니다. 비공개 id는 겹치면 안 됩니다.")
            continue
        s = seqs.get(pid)
        if s is None:
            errors.append(f"{i}행: id `{pid}`에 해당하는 FASTA 서열이 없습니다.")
            continue
        if not s:
            errors.append(f"{i}행: `{pid}`의 서열이 비어 있습니다.")
            continue
        badc = sorted(set(s) - ACGT)
        if badc:
            errors.append(f"{i}행: `{pid}`의 서열에 뉴클레오타이드가 아닌 문자가 있습니다: {''.join(badc[:8])}")
            continue
        if ref and config.MIN_COVERAGE_PCT and 100 * len(s) / ref[1] < config.MIN_COVERAGE_PCT:
            errors.append(f"{i}행: `{pid}`는 기준 게놈({ref[0]}, {ref[1]:,} bp)의 {100 * len(s) / ref[1]:.1f}%로 커버리지 하한 {config.MIN_COVERAGE_PCT:g}% 미만이라 큐레이션 대상이 아닙니다 "
                          "(하한은 GC_MIN_COVERAGE_PCT, 0이면 끔).")
            continue
        rec = dict(r)
        for k, v in defaults.items():
            if not rec.get(k):
                rec[k] = v
        d = (rec.get("date") or "").strip()
        if not rec.get("strain"):
            rec["strain"] = pid
            warnings.append(f"{i}행: `strain`이 비어 id로 채웠습니다.")
        if d:
            if re.fullmatch(r"\d{4}-(\d{2}|XX)-(\d{2}|XX)", d):
                rec["date"] = d
            else:
                nd = export.std_date(d)
                if nd:
                    rec["date"] = nd
                    if nd != d:
                        warnings.append(f"{i}행: 날짜 `{d}` → `{nd}`(Nextstrain 형식)로 바꿨습니다.")
                else:
                    rec["date"] = ""
                    warnings.append(f"{i}행: 날짜 `{d}`를 해석하지 못해 비웠습니다.")
        else:
            warnings.append(f"{i}행: 날짜가 비어 있습니다 (시간 기반 분석에서 제외될 수 있습니다).")
        if not rec.get("region") and rec.get("country"):
            rg = export.region_of(rec["country"])
            if rg:
                rec["region"] = rg
            else:
                warnings.append(f"{i}행: 국가 `{rec['country']}`의 대륙(region)을 알 수 없어 비웠습니다.")
        elif not rec.get("country"):
            warnings.append(f"{i}행: 국가가 비어 있습니다.")
        if rec.get("length") and str(rec["length"]).isdigit() and int(rec["length"]) != len(s):
            warnings.append(f"{i}행: `length` {rec['length']}과 서열 길이 {len(s):,}가 달라 서열 길이로 바꿨습니다.")
        rec["length"] = str(len(s))
        for k in list(rec):
            if k in RESERVED:
                rec.pop(k)
        rec["_id"], rec["_seq"] = pid, s
        out.append(rec)
    for pid in dup_seq:
        errors.append(f"FASTA에 id `{pid}`가 두 번 이상 있습니다.")
    extra = sorted(set(seqs) - set(seen))
    if extra:
        errors.append(f"메타데이터에 없는 FASTA 서열 {len(extra)}개: {', '.join(extra[:5])}{' …' if len(extra) > 5 else ''}")
    new = upd = same = 0
    for rec in out:
        old = have.get(rec["_id"])
        if not old:
            new += 1
            continue
        if _changes(json.loads(old["metadata_json"] or "{}"), old["seq_md5"], rec):
            upd += 1
        else:
            same += 1
    return {"ok": not errors, "errors": errors, "warnings": warnings, "rows": out, "n": len(rows), "new": new, "update": upd, "unchanged": same, "id_column": idc}


def _meta(rec: dict) -> dict:
    return {k: v for k, v in rec.items() if not k.startswith("_")}


def _changes(old_meta: dict, old_md5: str | None, rec: dict) -> dict:
    """What differs between the stored version and an upload: {field: [old, new]}, and 'sequence' when the sequence differs."""
    new_meta = _meta(rec)
    ch = {k: [old_meta.get(k, ""), new_meta.get(k, "")] for k in sorted(set(old_meta) | set(new_meta)) if str(old_meta.get(k, "")) != str(new_meta.get(k, ""))}
    if old_md5 != hashlib.md5(rec["_seq"].encode()).hexdigest():
        ch["sequence"] = [old_md5, hashlib.md5(rec["_seq"].encode()).hexdigest()]
    return ch


# ------------------------------------------------------------------------------------------------ import (every upload is a version)
def _write_fasta(path: Path, pid: str, seq: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f">{pid}\n" + "\n".join(seq[i:i + 70] for i in range(0, len(seq), 70)) + "\n", encoding="utf-8", newline="\n")


def _rel(p: Path) -> str:
    return p.relative_to(config.DATA_DIR).as_posix()


def _save_version(con, watch_id: int, d: Path, pid: str, rec: dict, source: str, note: str, now: str) -> tuple[str, int]:
    """Store one upload: the current files / row and a version record. -> (status 'new'|'update'|'unchanged', version)."""
    seq = rec["_seq"]
    meta = _meta(rec)
    md5 = hashlib.md5(seq.encode()).hexdigest()
    old = con.execute("SELECT * FROM private_genome WHERE watch_id=? AND id=?", (watch_id, pid)).fetchone()
    last = con.execute("SELECT MAX(version) FROM private_version WHERE watch_id=? AND id=?", (watch_id, pid)).fetchone()[0] or 0
    ch = _changes(json.loads(old["metadata_json"] or "{}"), old["seq_md5"], rec) if old else {}
    if old and not ch:
        return "unchanged", old["version"] or last
    ver = last + 1
    cur = d / f"{_safe(pid)}.fna"
    _write_fasta(cur, pid, seq)
    # a copy per version whose sequence is new (v1, or a changed sequence); other versions point at the copy they share
    prev_file = con.execute("SELECT file_path FROM private_version WHERE watch_id=? AND id=? ORDER BY version DESC LIMIT 1", (watch_id, pid)).fetchone()
    if not old or "sequence" in ch or not prev_file:
        vf = d / "_versions" / f"{_safe(pid)}.v{ver}.fna"
        _write_fasta(vf, pid, seq)
        vpath = _rel(vf)
    else:
        vpath = prev_file[0]
    con.execute("INSERT OR REPLACE INTO private_genome(watch_id,id,metadata_json,file_path,seq_md5,length,source_file,imported,updated,version) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (watch_id, pid, json.dumps(meta, ensure_ascii=False), _rel(cur), md5, len(seq), source, old["imported"] if old else now, now, ver))
    con.execute("INSERT INTO private_version(watch_id,id,version,metadata_json,file_path,seq_md5,length,source_file,uploaded,note,change_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (watch_id, pid, ver, json.dumps(meta, ensure_ascii=False), vpath, md5, len(seq), source, now, note,
                 json.dumps({"new": True} if not old else ch, ensure_ascii=False)))
    return ("update" if old else "new"), ver


def import_texts(watch_id: int, metadata, fasta_text: str, source_name: str = "", id_column: str | None = None, defaults: dict | None = None,
                 note: str = "") -> dict:
    """Import after a clean check. Errors stop everything (all or nothing). Each changed or new genome gets a new version."""
    rep = check(watch_id, metadata, fasta_text, id_column, defaults, source_name)
    if not rep["ok"]:
        raise ValueError("가져올 수 없습니다:\n" + "\n".join(rep["errors"][:20]))
    con = connect()
    try:
        w = con.execute("SELECT name FROM watchlist WHERE id=?", (watch_id,)).fetchone()
        d = private_dir(w["name"])
        now = datetime.now().isoformat(timespec="seconds")
        res = {"new": 0, "update": 0, "unchanged": 0}
        versions = {}
        for r in rep["rows"]:
            st, ver = _save_version(con, watch_id, d, r["_id"], r, source_name, note, now)
            res[st] += 1
            versions[r["_id"]] = {"status": st, "version": ver}
        audit(con, "private_import", f"{w['name']}: 신규 {res['new']}건, 새 버전 {res['update']}건, 변경 없음 {res['unchanged']}건 ({source_name})")
        con.commit()
    finally:
        con.close()
    _log.info("private import watch=%s %s", watch_id, res)
    return {**res, "warnings": rep["warnings"], "versions": versions}


def one(watch_id: int, fields: dict, sequence_text: str, source_name: str = "", dry_run: bool = False, defaults: dict | None = None, note: str = "") -> dict:
    """One genome: the fields of a single row and one sequence (FASTA text or bare sequence). Goes through the same check as a batch."""
    fields = {k.strip(): str(v).strip() for k, v in fields.items() if k and v is not None and str(v).strip() != ""}
    pid = fields.get("PPX_accession") or fields.get("accession") or ""
    fields.setdefault("PPX_accession", pid)
    seq = sequence_text.strip()
    if seq.startswith(">"):
        recs, _ = _read_fasta(seq)
        if len(recs) != 1:
            raise ValueError("서열은 한 개만 올릴 수 있습니다 (여러 개는 스프레드시트 업로드를 쓰세요).")
        seq = next(iter(recs.values()))
    seq = re.sub(r"\s+", "", seq).upper()
    cols = list(fields)
    meta_text = "\t".join(cols) + "\n" + "\t".join(fields[c] for c in cols) + "\n"
    fasta = f">{pid}\n{seq}\n"
    if dry_run:
        rep = check(watch_id, meta_text, fasta, "PPX_accession", defaults)
        return {k: rep[k] for k in ("ok", "errors", "warnings", "new", "update", "unchanged")}
    return import_texts(watch_id, meta_text, fasta, source_name or "(하나씩 입력)", "PPX_accession", defaults, note)


# ------------------------------------------------------------------------------------------------ listing, history, restore, delete
def list_rows(watch_id: int) -> list[dict]:
    con = connect()
    try:
        out = []
        qcs: dict[str, dict] = {}
        for q in con.execute("SELECT accession, tool, status, clade, lineage FROM qc_result WHERE watch_id=? AND accession LIKE 'private:%'", (watch_id,)):
            qcs.setdefault(q["accession"], {})[q["tool"]] = dict(q)
        from . import exclusions
        exm = exclusions.marked(watch_id, "private")
        for r in con.execute("SELECT id, metadata_json, length, imported, updated, version FROM private_genome WHERE watch_id=? ORDER BY id", (watch_id,)):
            m = json.loads(r["metadata_json"] or "{}")
            q = qcs.get(prefix(r["id"]), {})
            nc, st = q.get("nextclade") or {}, q.get("seqstats") or {}
            out.append({"id": r["id"], "strain": m.get("strain"), "date": m.get("date"), "country": m.get("country"), "clade": m.get("clade") or nc.get("clade") or "",
                        "length": r["length"], "imported": r["imported"], "updated": r["updated"], "version": r["version"] or 1,
                        "excluded": r["id"] in exm, "excl_reason": (exm.get(r["id"]) or {}).get("reason") or "",
                        "qc_stats": st.get("status", ""), "qc_nextclade": nc.get("status", ""), "nextclade_clade": nc.get("clade") or "",
                        "clade_differs": bool(m.get("clade") and nc.get("clade") and m["clade"] != nc["clade"])})
        return out
    finally:
        con.close()


def history(watch_id: int, pid: str | None = None) -> list[dict]:
    """Versions, newest first: {id, version, uploaded, source_file, note, changes (field list), sequence_changed, deleted, current}."""
    con = connect()
    try:
        cur = {r["id"]: r["version"] for r in con.execute("SELECT id, version FROM private_genome WHERE watch_id=?", (watch_id,))}
        q, a = "SELECT id, version, uploaded, source_file, note, change_json, length FROM private_version WHERE watch_id=?", [watch_id]
        if pid:
            q += " AND id=?"
            a.append(pid)
        out = []
        for r in con.execute(q + " ORDER BY id, version DESC", a):
            ch = json.loads(r["change_json"] or "{}")
            out.append({"id": r["id"], "version": r["version"], "uploaded": r["uploaded"], "source_file": r["source_file"], "note": r["note"] or "",
                        "new": bool(ch.get("new")), "deleted": bool(ch.get("deleted")), "sequence_changed": "sequence" in ch,
                        "changes": [k for k in ch if k not in ("new", "deleted", "sequence")], "detail": {k: v for k, v in ch.items() if k not in ("new", "deleted")},
                        "length": r["length"], "current": cur.get(r["id"]) == r["version"]})
        return out
    finally:
        con.close()


def restore(watch_id: int, pid: str, version: int) -> dict:
    """Bring an old version back as a NEW version (history is never rewritten)."""
    con = connect()
    try:
        v = con.execute("SELECT * FROM private_version WHERE watch_id=? AND id=? AND version=?", (watch_id, pid, version)).fetchone()
        if not v:
            raise ValueError("그 버전을 찾을 수 없습니다.")
        f = config.data_path(v["file_path"])
        if not f.exists():
            raise ValueError("그 버전의 서열 파일이 없습니다.")
        seqs, _ = _read_fasta(f.read_text(encoding="utf-8"))
        rec = dict(json.loads(v["metadata_json"] or "{}"))
        rec["_id"], rec["_seq"] = pid, next(iter(seqs.values()))
        w = con.execute("SELECT name FROM watchlist WHERE id=?", (watch_id,)).fetchone()
        st, ver = _save_version(con, watch_id, private_dir(w["name"]), pid, rec, v["source_file"] or "", f"버전 {version} 복원", datetime.now().isoformat(timespec="seconds"))
        audit(con, "private_restore", f"{pid}: v{version} -> v{ver}")
        con.commit()
    finally:
        con.close()
    return {"status": st, "version": ver}


def delete(watch_id: int, pid: str) -> bool:
    """Remove the current genome (row, current file, QC). Its versions stay and it can be restored."""
    con = connect()
    try:
        r = con.execute("SELECT * FROM private_genome WHERE watch_id=? AND id=?", (watch_id, pid)).fetchone()
        if not r:
            return False
        ver = (con.execute("SELECT MAX(version) FROM private_version WHERE watch_id=? AND id=?", (watch_id, pid)).fetchone()[0] or 0) + 1
        con.execute("INSERT INTO private_version(watch_id,id,version,metadata_json,file_path,seq_md5,length,source_file,uploaded,note,change_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (watch_id, pid, ver, r["metadata_json"], None, r["seq_md5"], r["length"], "", datetime.now().isoformat(timespec="seconds"), "삭제", json.dumps({"deleted": True})))
        con.execute("DELETE FROM private_genome WHERE watch_id=? AND id=?", (watch_id, pid))
        con.execute("DELETE FROM qc_result WHERE accession=?", (prefix(pid),))
        audit(con, "private_delete", f"{pid} (v{r['version'] or 1}까지의 이력은 남아 있어 복원할 수 있습니다)")
        con.commit()
    finally:
        con.close()
    try:
        config.data_path(r["file_path"]).unlink()
    except OSError:
        pass
    return True


def rows_for_export(watch_id: int) -> list[dict]:
    """[{id, meta, file_path, qc:{tool: {...}}}] for the export module."""
    con = connect()
    try:
        qcs: dict[str, dict] = {}
        for q in con.execute("SELECT accession, tool, status, clade, lineage, metrics_json FROM qc_result WHERE watch_id=? AND accession LIKE 'private:%'", (watch_id,)):
            qcs.setdefault(q["accession"], {})[q["tool"]] = {"status": q["status"], "clade": q["clade"], "lineage": q["lineage"], "m": json.loads(q["metrics_json"] or "{}")}
        return [{"id": r["id"], "meta": json.loads(r["metadata_json"] or "{}"), "file_path": r["file_path"], "qc": qcs.get(prefix(r["id"]), {}), "version": r["version"] or 1}
                for r in con.execute("SELECT id, metadata_json, file_path, version FROM private_genome WHERE watch_id=? ORDER BY id", (watch_id,))]
    finally:
        con.close()


# ------------------------------------------------------------------------------------------------ the spreadsheet template
def template_xlsx(virus: str | None = None, taxon: str | None = None) -> bytes:
    """A workbook to fill in: sheet `genomes` (header only - this is what is read), `열 설명` (every column) and `예시` (one filled row)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise ValueError("스프레드시트 템플릿을 만들려면 openpyxl이 필요합니다 (`mamba install openpyxl`).") from None
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    names = [c[0] for c in COLUMNS]
    ws.append(names)
    req = PatternFill("solid", fgColor="FDE9B8")
    for i, c in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        if c[1]:
            cell.fill = req
        ws.column_dimensions[cell.column_letter].width = max(14, len(c[0]) + 4)
    ws.freeze_panes = "A2"
    for r in range(2, 502):                                   # dates stay text: no silent re-formatting by the spreadsheet program
        ws.cell(row=r, column=names.index("date") + 1).number_format = "@"
    choices = CLADE_CHOICES.get((taxon or "").lower())
    if choices:
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "목록에 있는 clade만 고를 수 있습니다 (모르면 비워 두세요).", "clade"
        ws.add_data_validation(dv)
        col = ws.cell(row=2, column=names.index("clade") + 1).column_letter
        dv.add(f"{col}2:{col}500")
    g = wb.create_sheet("열 설명")
    g.append(["열", "필수", "설명", "예시"])
    for c in g[1]:
        c.font = Font(bold=True)
    for name, required, meaning, example, _ in COLUMNS:
        g.append([name, "필수" if required else "선택", meaning, example])
    g.append([])
    g.append(["참고", "", "`genomes` 시트의 첫 줄(머리글)과 아래 행만 읽습니다. 다른 열을 더 넣어도 보존됩니다(Nextstrain에 그대로 전달). FASTA는 따로 올리며 헤더는 id와 같아야 합니다.", ""])
    g.append(["", "", "region(대륙)·length·QC 열은 자동으로 채웁니다. 날짜는 2025-03-14 / 2025-03 / 2025 모두 가능합니다.", ""])
    g.column_dimensions["A"].width, g.column_dimensions["B"].width, g.column_dimensions["C"].width, g.column_dimensions["D"].width = 16, 8, 90, 34
    for row in g.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    e = wb.create_sheet("예시")
    e.append(names)
    e.append([c[3] for c in COLUMNS])
    for c in e[1]:
        c.font = Font(bold=True)
    if virus:
        e["A4"] = f"{virus} — 이 시트는 읽지 않습니다."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_tsv() -> str:
    return "\t".join(c[0] for c in COLUMNS) + "\n"
